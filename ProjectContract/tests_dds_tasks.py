# -*- coding: utf-8 -*-
"""Колонка задач в ДДС-отчёте: связки EstimateTaskLink + основная задача,
фильтры «с задачами» / «без задач», колонка в xlsx."""
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse

from ProjectContract.models import (
    Act, Contract, ContractEstimate, Contractor, EstimateConcept,
    EstimateTaskLink,
)
from ProjectTDL.models import TaskNode
from StaticData.models import (
    BuildingNumber, BuildingType, Category, DesignChapter, ProjectSite,
    Status,
)


class DdsTasksColumnTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import User
        cls.owner = User.objects.create_superuser(
            username='dds_owner', password='pw', email='d@d.d')
        cls.site = ProjectSite.objects.create(name='DT-Объект')
        cls.contractor = Contractor.objects.create(name='DT-Подрядчик')
        cls.contract = Contract.objects.create(
            project_site=cls.site, contractor=cls.contractor,
            name='DT-Договор', price=Decimal('1000000.00'))
        cls.estimate = ContractEstimate.objects.create(
            contract=cls.contract, name='DT-Смета')
        Status.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        Category.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        cls.bare = EstimateConcept.objects.create(
            estimate=cls.estimate, name='Поз-без-задач',
            amount=Decimal('1000.00'))
        cls.with_primary = EstimateConcept.objects.create(
            estimate=cls.estimate, name='Поз-с-основной',
            amount=Decimal('2000.00'))
        cls.task_primary = TaskNode.objects.create(
            owner=cls.owner, name='Главная задача', project_site=cls.site)
        cls.with_primary.task_node = cls.task_primary
        cls.with_primary.save(update_fields=['task_node'])
        cls.with_links = EstimateConcept.objects.create(
            estimate=cls.estimate, name='Поз-со-связками',
            amount=Decimal('3000.00'))
        cls.task_a = TaskNode.objects.create(
            owner=cls.owner, name='Задача А', project_site=cls.site)
        cls.task_b = TaskNode.objects.create(
            owner=cls.owner, name='Задача Б', project_site=cls.site)
        EstimateTaskLink.objects.create(
            estimate_concept=cls.with_links, task_node=cls.task_a)
        EstimateTaskLink.objects.create(
            estimate_concept=cls.with_links, task_node=cls.task_b)

    def setUp(self):
        self.client.force_login(self.owner)

    def test_column_shows_tasks(self):
        response = self.client.get(reverse('dds_report'))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode('utf-8')
        self.assertIn('Задачи', html)
        self.assertIn('Главная задача', html)
        self.assertIn('Задача А', html)
        self.assertIn('Задача Б', html)

    def test_filter_without_tasks(self):
        response = self.client.get(reverse('dds_report'), {'tasks': 'none'})
        rows = list(response.context['rows'])
        self.assertEqual([c.pk for c in rows], [self.bare.pk])

    def test_filter_with_tasks(self):
        response = self.client.get(reverse('dds_report'), {'tasks': 'has'})
        pks = {c.pk for c in response.context['rows']}
        self.assertEqual(pks, {self.with_primary.pk, self.with_links.pk})

    def test_xlsx_has_tasks_column(self):
        import openpyxl
        from io import BytesIO
        response = self.client.get(reverse('dds_report'), {'export': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(filename=BytesIO(response.content), data_only=True)
        ws = wb['ДДС']
        headers = [c.value for c in ws[1]]
        self.assertIn('Задачи', headers)
        col = headers.index('Задачи') + 1
        names = [ws.cell(row=r, column=col).value or '' for r in range(2, ws.max_row - 1)]
        blob = ' '.join(names)
        self.assertIn('Главная задача', blob)
        self.assertIn('Задача А', blob)
