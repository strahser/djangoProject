# -*- coding: utf-8 -*-
"""P2: ДДС-отчёт ProjectContract (вынесено из views.py 1:1)."""
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import Prefetch, Q
from django.http import HttpResponse
from django.shortcuts import render

from ProjectContract.models import (
    Act, Contractor, ContractEstimate, EstimateConcept, EstimateTaskLink,
)
from StaticData.models import BuildingNumber, DesignChapter

from .models import Contract


@login_required
def dds_report(request):
    """Отчёт по позициям смет в формате ДДС: Здание | Раздел | Наименование |
    Стоимость | Договор | Закрытие (акт) | Дата акта | Задачи.

    Цепочка Договор → Смета → Позиция → Акт, всё ссылками. Фильтры:
    подрядчик / договор / здание / раздел / закрытие / задачи.
    Выгрузка ?export=xlsx.
    """
    sel_contractor = request.GET.get('contractor') or ''
    sel_contract = request.GET.get('contract') or ''
    sel_building = request.GET.get('building') or ''
    sel_chapter = request.GET.get('chapter') or ''
    sel_acts = [v for v in request.GET.getlist('act') if v]
    sel_tasks = request.GET.get('tasks') or ''

    qs = EstimateConcept.objects.select_related(
        'estimate__contract__contractor', 'building_number__name',
        'design_chapter', 'act', 'estimate', 'task_node',
    ).prefetch_related(
        Prefetch('task_links', queryset=EstimateTaskLink.objects.select_related('task_node')),
    )
    if sel_contractor:
        qs = qs.filter(estimate__contract__contractor_id=sel_contractor)
    if sel_contract:
        qs = qs.filter(estimate__contract_id=sel_contract)
    if sel_building:
        qs = qs.filter(building_number_id=sel_building)
    if sel_chapter:
        qs = qs.filter(design_chapter_id=sel_chapter)
    if sel_acts:
        act_q = Q()
        act_ids = [v for v in sel_acts if v != '__none__']
        if act_ids:
            act_q |= Q(act_id__in=act_ids)
        if '__none__' in sel_acts:
            act_q |= Q(act__isnull=True)
        qs = qs.filter(act_q)
    if sel_tasks == 'none':
        qs = qs.filter(task_links__isnull=True, task_node__isnull=True)
    elif sel_tasks == 'has':
        qs = qs.filter(Q(task_links__isnull=False) | Q(task_node__isnull=False)).distinct()
    qs = qs.order_by('estimate__contract__name', 'building_number__name__name',
                     'building_number__building_number', 'id')

    rows = list(qs)
    for c in rows:
        # Задачи позиции: основная первой, дальше связки (без дублей).
        seen, tasks = set(), []
        if c.task_node_id and c.task_node is not None:
            seen.add(c.task_node_id)
            tasks.append(c.task_node)
        for link in c.task_links.all():
            t = link.task_node
            if t is not None and t.pk not in seen:
                seen.add(t.pk)
                tasks.append(t)
        c.task_list = tasks
    total = sum((c.amount or 0 for c in rows), Decimal('0'))
    closed_total = sum((c.amount or 0 for c in rows if c.act_id), Decimal('0'))

    contracts = Contract.objects.order_by('name')
    if sel_contractor:
        contracts = contracts.filter(contractor_id=sel_contractor)
    acts = Act.objects.select_related('contract').order_by(
        'contract__name', 'act_date', 'number')
    filt = {'sel_contractor': sel_contractor, 'sel_contract': sel_contract,
            'sel_building': sel_building, 'sel_chapter': sel_chapter,
            'sel_tasks': sel_tasks,
            'sel_acts': sel_acts, 'act_summary': _act_filter_summary(sel_acts)}

    if request.GET.get('export') == 'xlsx':
        return _dds_report_xlsx(rows, total, closed_total)

    return render(request, 'ProjectContract/dds_report.html', {
        'rows': rows, 'total': total, 'closed_total': closed_total,
        'open_total': total - closed_total,
        'contracts': contracts,
        'contractors': Contractor.objects.order_by('name'),
        'buildings': BuildingNumber.objects.select_related('name').order_by('name__name', 'building_number'),
        'chapters': DesignChapter.objects.order_by('short_name'),
        'acts': acts,
        'sel': filt,
    })




def _act_filter_summary(sel_acts):
    """Краткий итог для кнопки фильтра актов."""
    if not sel_acts:
        return 'Все'
    vals = [v for v in sel_acts if v != '__none__']
    if not vals:
        return 'Без акта'
    if len(sel_acts) == 1:
        try:
            return str(Act.objects.select_related('contract').get(pk=vals[0]))
        except (Act.DoesNotExist, ValueError):
            return 'Все'
    suffix = ' + без акта' if '__none__' in sel_acts else ''
    return f'Выбрано: {len(vals)}{suffix}'




def _dds_report_xlsx(rows, total, closed_total):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from django.utils import timezone

    wb = Workbook()
    ws = wb.active
    ws.title = 'ДДС'
    headers = ['Здание', 'Раздел', 'Наименование', 'Стоимость',
               'Договор', 'Закрытие (акт)', 'Дата акта', 'Задачи']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for c in rows:
        b = c.building_number
        ws.append([
            str(b) if b else '',
            c.design_chapter.name if c.design_chapter else '',
            c.name or '',
            float(c.amount or 0),
            c.estimate.contract.name if c.estimate and c.estimate.contract else '',
            f'Акт №{c.act.number}' if c.act else '',
            c.act.act_date if c.act and c.act.act_date else '',
            '; '.join(t.name for t in getattr(c, 'task_list', [])),
        ])
    ws.append(['', '', 'ИТОГО:', float(total), '', '', '', ''])
    ws.append(['', '', 'в т.ч. закрыто актами:', float(closed_total), '', '', '', ''])
    for cell in ws[ws.max_row - 1]:
        cell.font = Font(bold=True)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    widths = [28, 28, 70, 14, 30, 18, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:H{ws.max_row}'
    ws.freeze_panes = 'A2'
    stamp = timezone.now().strftime('%Y%m%d_%H%M')
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="DDS_{stamp}.xlsx"'
    wb.save(response)
    return response
