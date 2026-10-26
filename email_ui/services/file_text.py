# -*- coding: utf-8 -*-
"""Текст вложений для полнотекстового поиска (pdf/docx/xlsx/txt).

Тяжёлые/неподдержанные форматы — пусто (не ошибка). Несуществующий файл —
исключение FileNotFoundError, чтобы вызыватель мог отложить индексацию
(файлы лежат на сетевом диске, которого может не быть).
"""
import os
import re

MAX_TEXT_CHARS = 200_000
MAX_FILE_BYTES = 30 * 1024 * 1024
TEXT_EXTENSIONS = frozenset(
    {'txt', 'csv', 'log', 'md', 'eml', 'xml', 'json', 'html', 'htm'})
_WS_RE = re.compile(r'\s+')


def _collapse(text):
    return _WS_RE.sub(' ', text or '').strip()


def _from_txt(path):
    for enc in ('utf-8', 'windows-1251'):
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeError, ValueError):
            continue
    return ''


def _from_pdf(path, budget):
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    total = 0
    for page in reader.pages:
        try:
            text = page.extract_text() or ''
        except Exception:
            continue
        if text:
            parts.append(text)
            total += len(text)
            if total >= budget:
                break
    return '\n'.join(parts)


def _from_docx(path, budget):
    from docx import Document
    doc = Document(path)
    parts = []
    total = 0
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
            total += len(p.text)
            if total >= budget:
                break
    if total < budget:
        for table in doc.tables:
            for row in table.rows:
                line = ' '.join(c.text for c in row.cells)
                if line.strip():
                    parts.append(line)
                    total += len(line)
                    if total >= budget:
                        break
            if total >= budget:
                break
    return '\n'.join(parts)


def _from_xlsx(path, budget):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            line = ' '.join(str(v) for v in row if v is not None and str(v).strip())
            if line:
                parts.append(line)
                total += len(line)
                if total >= budget:
                    break
        if total >= budget:
            break
    wb.close()
    return '\n'.join(parts)


def extract_file_text(path, filename='', content_type='', max_chars=MAX_TEXT_CHARS):
    """Текст файла для поиска ('' если извлечь нельзя)."""
    if not path or not os.path.isfile(path):
        raise FileNotFoundError(path)
    try:
        if os.path.getsize(path) > MAX_FILE_BYTES:
            return ''
    except OSError:
        raise FileNotFoundError(path)
    ext = (os.path.splitext(filename or path)[1].lower().lstrip('.')
           or 'unknown')
    ctype = (content_type or '').lower()
    try:
        if ext == 'pdf' or 'pdf' in ctype:
            text = _from_pdf(path, max_chars)
        elif ext == 'docx' or 'wordprocessingml' in ctype:
            text = _from_docx(path, max_chars)
        elif ext == 'xlsx' or 'spreadsheetml' in ctype:
            text = _from_xlsx(path, max_chars)
        elif ext in TEXT_EXTENSIONS or ctype.startswith('text/'):
            text = _from_txt(path)
        else:
            return ''
    except Exception:
        return ''
    return _collapse(text)[:max_chars]


def index_attachment(att, max_chars=MAX_TEXT_CHARS):
    """Извлечь текст вложения и сохранить флаг. True — проиндексировано
    (пусть даже пусто); файл отсутствует — FileNotFoundError (не помечаем)."""
    text = extract_file_text(att.resolved_path, att.filename,
                             getattr(att, 'content_type', ''),
                             max_chars=max_chars)
    att.text_content = text
    att.text_indexed = True
    att.save(update_fields=['text_content', 'text_indexed'])
    return True
