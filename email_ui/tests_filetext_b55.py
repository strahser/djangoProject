# -*- coding: utf-8 -*-
"""Поиск по тексту вложений: извлечение, индексация, фильтр «В файлах»."""
import os
import tempfile

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase

from Emails.models import Attachment, Email
from email_ui.services import file_text
from email_ui.services.file_text import extract_file_text, index_attachment
from email_ui.services.query_service import filter_emails
from StaticData.models import Category, Status

MINIMAL_PDF = (
    b'%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n'
    b'2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n'
    b'3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj\n'
    b'4 0 obj<</Length 58>>stream\nBT /F1 24 Tf 100 700 Td (Hello PDF world) Tj ET\nendstream\nendobj\n'
    b'5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n'
    b'xref\n0 6\n0000000000 65535 f \n0000000009 00000 n \n0000000056 00000 n \n'
    b'0000000111 00000 n \n0000000262 00000 n \n0000000370 00000 n \n'
    b'trailer<</Size 6/Root 1 0 R>>\nstartxref\n434\n%%EOF'
)


def _tmp(suffix, data, mode='wb'):
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, mode) as f:
        f.write(data)
    return path


class ExtractFileTextTest(TestCase):
    def test_txt_encodings(self):
        p = _tmp('.txt', 'Смета Изм 5'.encode('windows-1251'))
        self.assertIn('Смета', extract_file_text(p, 'a.txt'))

    def test_pdf(self):
        p = _tmp('.pdf', MINIMAL_PDF)
        self.assertIn('Hello PDF world', extract_file_text(p, 'a.pdf'))

    def test_docx(self):
        from docx import Document
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        doc = Document()
        doc.add_paragraph('Ведомость работ Коровник')
        doc.save(path)
        self.assertIn('Коровник', extract_file_text(path, 'a.docx'))

    def test_xlsx(self):
        from openpyxl import Workbook
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb = Workbook()
        wb.active['A1'] = 'Спецификация насосов'
        wb.save(path)
        wb.close()
        self.assertIn('Спецификация', extract_file_text(path, 'a.xlsx'))

    def test_unsupported_returns_empty(self):
        p = _tmp('.exe', b'MZ\x90\x00binary')
        self.assertEqual(extract_file_text(p, 'a.exe'), '')

    def test_missing_raises(self):
        with self.assertRaises(FileNotFoundError):
            extract_file_text('/nonexistent/file.pdf', 'file.pdf')


class IndexAttachmentsTest(TestCase):
    def setUp(self):
        Category.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        Status.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        self.mail = Email.objects.create(
            uid='f-idx-1', email_type='IN', subject='S', sender='a@t.co',
            folder='inbox', sent_status='sent')
        self.path = _tmp('.txt', 'Изыскания по площадке'.encode('utf-8'))
        self.att = Attachment.objects.create(
            email=self.mail, filename='notes.txt', file_path=self.path,
            size=10, content_type='text/plain')

    def test_index_attachment_saves_text_and_flag(self):
        self.assertTrue(index_attachment(self.att))
        self.att.refresh_from_db()
        self.assertTrue(self.att.text_indexed)
        self.assertIn('Изыскания', self.att.text_content)

    def test_index_missing_file_leaves_unindexed(self):
        self.att.file_path = '/nonexistent/x.txt'
        with self.assertRaises(FileNotFoundError):
            index_attachment(self.att)
        self.att.refresh_from_db()
        self.assertFalse(self.att.text_indexed)

    def test_command_indexes_unindexed_only(self):
        call_command('index_attachments')
        self.att.refresh_from_db()
        self.assertTrue(self.att.text_indexed)
        missing = Attachment.objects.create(
            email=self.mail, filename='gone.pdf', file_path='/nonexistent/gone.pdf',
            size=1, content_type='application/pdf')
        call_command('index_attachments')
        missing.refresh_from_db()
        self.assertFalse(missing.text_indexed)


class FileSearchFilterTest(TestCase):
    def setUp(self):
        Category.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        Status.objects.get_or_create(pk=1, defaults={'name': 'Default'})
        self.m1 = Email.objects.create(
            uid='f-s-1', email_type='IN', subject='Письмо без совпадений',
            sender='a@t.co', folder='inbox', sent_status='sent')
        self.m2 = Email.objects.create(
            uid='f-s-2', email_type='IN', subject='Другое письмо',
            sender='b@t.co', folder='inbox', sent_status='sent')
        path = _tmp('.txt', 'СМЕТА на монтаж вентиляции'.encode('utf-8'))
        self.att = Attachment.objects.create(
            email=self.m2, filename='s.txt', file_path=path,
            size=10, content_type='text/plain')
        index_attachment(self.att)

    def _pks(self, **params):
        params.setdefault('folder', 'inbox')
        return set(filter_emails(
            Email.objects.filter(folder='inbox'), params
        ).values_list('pk', flat=True))

    def test_without_in_files_not_found(self):
        self.assertNotIn(self.m2.pk, self._pks(search='вентиляции'))

    def test_with_in_files_found_cyrillic_case(self):
        pks = self._pks(search='смета', in_files=True)
        self.assertIn(self.m2.pk, pks)
        self.assertNotIn(self.m1.pk, pks)

    def test_in_files_no_duplicates(self):
        Attachment.objects.create(
            email=self.m2, filename='s2.txt', file_path=self.att.file_path,
            size=10, content_type='text/plain', text_content='смета',
            text_indexed=True)
        pks = list(filter_emails(
            Email.objects.filter(folder='inbox'),
            {'folder': 'inbox', 'search': 'смета', 'in_files': True},
        ).values_list('pk', flat=True))
        self.assertEqual(pks.count(self.m2.pk), 1)
